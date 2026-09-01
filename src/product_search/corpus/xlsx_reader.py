from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from product_search.corpus.rules import RejectionReason, validate_product_id


class WorkbookImportError(ValueError):
    """XLSX нельзя использовать как подготовленный снимок корпуса."""


@dataclass(frozen=True)
class ParsedProduct:
    product_id: str
    source_row: int
    category: str
    primary_image: bytes


@dataclass(frozen=True)
class RejectedRow:
    source_row: int
    reason: RejectionReason


@dataclass(frozen=True)
class WorkbookInspection:
    accepted: tuple[ParsedProduct, ...]
    rejected: tuple[RejectedRow, ...]
    duplicate_product_ids: tuple[str, ...]

    @property
    def has_critical_error(self) -> bool:
        return bool(self.duplicate_product_ids)


@dataclass(frozen=True)
class WorkbookLayout:
    sheet_name: str | None
    header_row: int
    first_data_row: int
    product_id_header: str
    category_header: str
    primary_image_header: str


SYNTHETIC_LAYOUT = WorkbookLayout(
    sheet_name=None,
    header_row=1,
    first_data_row=2,
    product_id_header="product_id",
    category_header="category",
    primary_image_header="main_image",
)

PRODUCT_NOVELTIES_LAYOUT = WorkbookLayout(
    sheet_name="На рассм-е",
    header_row=2,
    first_data_row=6,
    product_id_header="product_id",
    category_header="Категория",
    primary_image_header="Фотография",
)


def inspect_workbook(
    path: Path, layout: WorkbookLayout = SYNTHETIC_LAYOUT
) -> WorkbookInspection:
    if path.suffix.lower() != ".xlsx":
        raise WorkbookImportError("Служебная загрузка принимает только файл XLSX.")
    try:
        workbook = load_workbook(path, read_only=False, data_only=True)
    except Exception as error:
        raise WorkbookImportError("XLSX повреждён или не читается.") from error

    try:
        worksheet = workbook[layout.sheet_name] if layout.sheet_name else workbook.active
    except KeyError as error:
        raise WorkbookImportError("В XLSX отсутствует лист подготовленного снимка.") from error
    columns = _header_columns(worksheet, layout)
    images = _images_by_cell(worksheet)
    accepted: list[ParsedProduct] = []
    rejected: list[RejectedRow] = []
    seen_ids: set[str] = set()
    duplicates: set[str] = set()

    for row in range(layout.first_data_row, worksheet.max_row + 1):
        raw_id = worksheet.cell(row, columns["product_id"]).value
        raw_category = worksheet.cell(row, columns["category"]).value
        if raw_id is None and raw_category is None:
            continue
        product_id = str(raw_id).strip() if raw_id is not None else ""
        reason = validate_product_id(product_id)
        if reason is not None:
            rejected.append(RejectedRow(row, reason))
            continue
        if product_id in seen_ids:
            duplicates.add(product_id)
        seen_ids.add(product_id)
        category = str(raw_category).strip() if raw_category is not None else ""
        if not category:
            rejected.append(RejectedRow(row, RejectionReason.MISSING_CATEGORY))
            continue
        primary_image = images.get((row, columns["primary_image"]))
        if primary_image is None:
            rejected.append(RejectedRow(row, RejectionReason.MISSING_PRIMARY_IMAGE))
            continue
        accepted.append(ParsedProduct(product_id, row, category, primary_image))

    return WorkbookInspection(tuple(accepted), tuple(rejected), tuple(sorted(duplicates)))


def _header_columns(worksheet, layout: WorkbookLayout) -> dict[str, int]:
    columns = {
        str(cell.value).strip(): cell.column
        for cell in worksheet[layout.header_row]
        if cell.value is not None and str(cell.value).strip()
    }
    required_columns = {
        "product_id": layout.product_id_header,
        "category": layout.category_header,
        "primary_image": layout.primary_image_header,
    }
    missing = [name for name in required_columns.values() if name not in columns]
    if missing:
        raise WorkbookImportError("В XLSX отсутствуют обязательные колонки подготовленного снимка.")
    return {key: columns[value] for key, value in required_columns.items()}


def _images_by_cell(worksheet) -> dict[tuple[int, int], bytes]:
    images: dict[tuple[int, int], bytes] = {}
    for image in worksheet._images:
        anchor = image.anchor._from
        images[(anchor.row + 1, anchor.col + 1)] = image._data()
    return images
