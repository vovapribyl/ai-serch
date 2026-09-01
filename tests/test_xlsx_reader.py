from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlsxImage
from PIL import Image

from product_search.corpus.rules import RejectionReason
from product_search.corpus.xlsx_reader import (
    PRODUCT_NOVELTIES_LAYOUT,
    WorkbookImportError,
    inspect_workbook,
)


def write_fixture(path: Path, rows: list[tuple[object, object]], image_rows: set[int]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["product_id", "category", "main_image", "secondary_image"])
    for product_id, category in rows:
        sheet.append([product_id, category])
    for row in image_rows:
        image = Image.new("RGB", (1, 1), "red")
        buffer = BytesIO()
        image.save(buffer, format="PNG")
        buffer.seek(0)
        xlsx_image = XlsxImage(buffer)
        xlsx_image.anchor = f"C{row}"
        sheet.add_image(xlsx_image)
    workbook.save(path)


def test_reads_primary_image_and_excludes_missing_fields(tmp_path: Path) -> None:
    path = tmp_path / "fixture.xlsx"
    write_fixture(path, [("101", "Накладные ногти"), ("102", ""), ("103", "Накладные ногти")], {2, 3})

    result = inspect_workbook(path)

    assert [record.product_id for record in result.accepted] == ["101"]
    assert [(row.source_row, row.reason) for row in result.rejected] == [
        (3, RejectionReason.MISSING_CATEGORY),
        (4, RejectionReason.MISSING_PRIMARY_IMAGE),
    ]


def test_duplicate_product_id_is_critical_and_does_not_depend_on_image(tmp_path: Path) -> None:
    path = tmp_path / "fixture.xlsx"
    write_fixture(path, [("101", "Накладные ногти"), ("101", "Накладные ногти")], {2, 3})

    result = inspect_workbook(path)

    assert result.duplicate_product_ids == ("101",)
    assert result.has_critical_error


def test_rejects_non_xlsx_and_unreadable_workbook(tmp_path: Path) -> None:
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"not a workbook")

    with pytest.raises(WorkbookImportError, match="повреждён"):
        inspect_workbook(path)


def test_reads_product_novelties_layout_with_russian_headers(tmp_path: Path) -> None:
    path = tmp_path / "product-novelties.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "На рассм-е"
    sheet.cell(2, 6, "Категория")
    sheet.cell(2, 13, "Фотография")
    sheet.cell(2, 87, "product_id")
    sheet.cell(6, 6, "Накладные ногти")
    sheet.cell(6, 87, 1000006)
    image = Image.new("RGB", (1, 1), "red")
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)
    xlsx_image = XlsxImage(buffer)
    xlsx_image.anchor = "M6"
    sheet.add_image(xlsx_image)
    workbook.save(path)

    result = inspect_workbook(path, layout=PRODUCT_NOVELTIES_LAYOUT)

    assert [(record.product_id, record.source_row, record.category) for record in result.accepted] == [
        ("1000006", 6, "Накладные ногти")
    ]
