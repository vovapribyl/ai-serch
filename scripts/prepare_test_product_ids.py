"""Create a non-destructive test copy of an XLSX workbook with product IDs."""

from __future__ import annotations

import argparse
from copy import copy
from pathlib import Path

from openpyxl import load_workbook


def is_present(value: object) -> bool:
    return value is not None and str(value).strip() != ""


def copy_cell_style(source: object, target: object) -> None:
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy(source.protection)


def prepare_copy(
    input_path: Path,
    output_path: Path,
    sheet_name: str,
    header_row: int = 2,
    first_data_row: int = 6,
    category_column: int = 6,
) -> tuple[int, int]:
    workbook = load_workbook(input_path, data_only=False, keep_links=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Sheet not found: {sheet_name}")

    worksheet = workbook[sheet_name]
    target_column = worksheet.max_column + 1
    header_cell = worksheet.cell(header_row, target_column, "product_id")
    copy_cell_style(worksheet.cell(header_row, category_column), header_cell)

    assigned_count = 0
    for row_number in range(first_data_row, worksheet.max_row + 1):
        if not is_present(worksheet.cell(row_number, category_column).value):
            continue
        target_cell = worksheet.cell(row_number, target_column, 1_000_000 + row_number)
        copy_cell_style(worksheet.cell(row_number, category_column), target_cell)
        target_cell.number_format = "0"
        assigned_count += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return target_column, assigned_count


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--sheet", required=True)
    arguments = parser.parse_args()

    target_column, assigned_count = prepare_copy(
        input_path=arguments.input,
        output_path=arguments.output,
        sheet_name=arguments.sheet,
    )
    print(
        f"sheet={arguments.sheet}; target_column={target_column}; "
        f"assigned_product_ids={assigned_count}"
    )


if __name__ == "__main__":
    main()
